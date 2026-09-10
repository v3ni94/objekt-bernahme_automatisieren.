"""Protokoll des Ordnerabgleichs (Fachentwurf F 10): JSON je Lauf, Excel je Lauf und Sammelfassung ueber alle Objekte.
Ablage unter /srv/objektakte/exports/drive-sync/ (B-14). Die Datenbank bleibt Quelle der Wahrheit."""

from __future__ import annotations

import json
from pathlib import Path

from django.conf import settings
from django.utils import timezone

from apps.drive.models import DriveSyncAction, DriveSyncRun
from apps.objects.models import ManagedObject
from apps.review.models import ReviewCase


def export_dir() -> Path:
    path = Path(settings.OBJEKTAKTE["DATA_DIR"]) / "exports" / "drive-sync"
    path.mkdir(parents=True, exist_ok=True)
    return path


def run_to_dict(run: DriveSyncRun) -> dict:
    obj = run.object
    summary = run.summary or {}
    actions = [
        {
            "seq_no": a.seq_no,
            "action_type": a.action_type,
            "category": summary.get("categories", {}).get(str(a.seq_no)),
            "name_before": a.name_before,
            "name_after": a.name_after,
            "drive_id": a.target_drive_id,
            "parent_id": a.parent_drive_id,
            "file_count_before": a.file_count_before,
            "file_count_after": a.file_count_after,
            "planned": a.planned,
            "executed": a.executed,
            "result": a.result,
            "error": a.error_message,
            "executed_at": a.executed_at.isoformat() if a.executed_at else None,
            "note": summary.get("notes", {}).get(str(a.seq_no)),
        }
        for a in DriveSyncAction.objects.filter(sync_run=run).order_by("seq_no")
    ]
    reviews = (
        list(
            ReviewCase.objects.filter(context__contains={"sync_run_id": run.pk}).values(
                "id", "case_type", "case_subtype", "status"
            )
        )
        if obj
        else []
    )
    return {
        "kopf": {
            "objektnummer": obj.object_number if obj else None,
            "objekt": obj.name if obj else None,
            "verwaltungsart": obj.management_type if obj else None,
            "lauf_id": run.pk,
            "modus": "Dry-Run" if run.dry_run else "Ausführung",
            "ausloeser": summary.get("trigger"),
            "nutzer": run.triggered_by.email if run.triggered_by else None,
            "beginn": run.started_at.isoformat(),
            "ende": run.finished_at.isoformat() if run.finished_at else None,
            "status": run.status,
            "objektordner_id": obj.drive_root_folder_id if obj else None,
            "objektordner_name": obj.drive_root_folder_name if obj else None,
            "root_matches": run.root_matches,
            "dateien_vorher": run.file_count_before,
            "dateien_nachher": run.file_count_after,
            "id_hash_vorher": summary.get("id_hash_before"),
            "id_hash_nachher": summary.get("id_hash_after"),
            "aktionen_geplant": run.actions_planned,
            "aktionen_ausgefuehrt": run.actions_executed,
            "no_changes": run.no_changes,
            "fehler": run.error_message,
        },
        "aktionen": actions,
        "hinweise": summary.get("hints", []),
        "ignorierte_wurzelordner": summary.get("ignored_root_folders", []),
        "review_faelle": reviews,
        "inventur": summary.get("inventory", {}),
        "nachbedingungen_ok": summary.get("postconditions_ok"),
    }


def _filename(run: DriveSyncRun, ext: str) -> str:
    number = run.object.object_number if run.object else "alle"
    return f"Abgleichsprotokoll_{number}_{timezone.localtime(run.started_at):%d.%m.%Y}_{run.pk}.{ext}"


def write_json(run: DriveSyncRun) -> Path:
    path = export_dir() / _filename(run, "json")
    path.write_text(json.dumps(run_to_dict(run), ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return path


def write_xlsx(run: DriveSyncRun) -> Path:
    import openpyxl
    from openpyxl.styles import Font

    data = run_to_dict(run)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kopf"
    for k, v in data["kopf"].items():
        ws.append([k, v if not isinstance(v, list | dict) else json.dumps(v, ensure_ascii=False)])
    ws2 = wb.create_sheet("Aktionen")
    cols = [
        "seq_no",
        "action_type",
        "category",
        "name_before",
        "name_after",
        "drive_id",
        "parent_id",
        "file_count_before",
        "file_count_after",
        "planned",
        "executed",
        "result",
        "error",
        "executed_at",
        "note",
    ]
    ws2.append(cols)
    for c in ws2[1]:
        c.font = Font(bold=True)
    for a in data["aktionen"]:
        ws2.append([a.get(c) for c in cols])
    ws3 = wb.create_sheet("Hinweise")
    for h in data["hinweise"]:
        ws3.append([h])
    for h in data["ignorierte_wurzelordner"]:
        ws3.append([f"ignorierter Wurzelordner: {h}"])
    ws4 = wb.create_sheet("Review-Fälle")
    ws4.append(["id", "case_type", "case_subtype", "status"])
    for r in data["review_faelle"]:
        ws4.append([r["id"], r["case_type"], r["case_subtype"], r["status"]])
    ws5 = wb.create_sheet("Inventur")
    for k, v in (data["inventur"] or {}).items():
        ws5.append([k, v])
    path = export_dir() / _filename(run, "xlsx")
    wb.save(path)
    return path


def write_summary_xlsx() -> Path:
    """Sammelfassung ueber alle Objekte mit letztem Lauf (F 10.3): Zeilen mit offenen Review-Faellen gelb."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Objekte"
    ws.append(
        [
            "Objekt",
            "Bezeichnung",
            "Objektordner-ID",
            "letzter Lauf",
            "Modus",
            "Status",
            "Dateien vorher",
            "Dateien nachher",
            "Umbenennungen",
            "Anlagen",
            "offene Review-Fälle",
            "no_changes",
        ]
    )
    for c in ws[1]:
        c.font = Font(bold=True)
    yellow = PatternFill(start_color="FFF7E8", end_color="FFF7E8", fill_type="solid")
    for obj in ManagedObject.active.order_by("object_number_numeric"):
        run = DriveSyncRun.objects.filter(object=obj).order_by("-started_at").first()
        open_cases = ReviewCase.objects.filter(
            object=obj,
            case_type__in=["drive_structure", "duplicate_object_number"],
            status__in=["open", "in_progress"],
        ).count()
        if run is None:
            ws.append(
                [
                    obj.object_number,
                    obj.name,
                    obj.drive_root_folder_id,
                    None,
                    None,
                    "kein Lauf",
                    None,
                    None,
                    None,
                    None,
                    open_cases,
                    None,
                ]
            )
            continue
        acts = DriveSyncAction.objects.filter(sync_run=run)
        ws.append(
            [
                obj.object_number,
                obj.name,
                obj.drive_root_folder_id,
                timezone.localtime(run.started_at).strftime("%d.%m.%Y %H:%M"),
                "Dry-Run" if run.dry_run else "Ausführung",
                run.status,
                run.file_count_before,
                run.file_count_after,
                acts.filter(action_type="rename_folder", result="ok").count(),
                acts.filter(action_type="create_folder", result="ok").count(),
                open_cases,
                run.no_changes,
            ]
        )
        if open_cases:
            for c in ws[ws.max_row]:
                c.fill = yellow
    path = export_dir() / f"Abgleichsprotokoll_Bestand_{timezone.localdate():%d.%m.%Y}.xlsx"
    wb.save(path)
    return path
