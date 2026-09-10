"""Excel-Ausgabe der Listen (H 5.3, openpyxl): Blaetter Aktuell, Historie, Offene Punkte (optional Herkunft),
Dokumentkopf in den Zeilen 1 bis 3, Tabellenobjekt mit Filter ab Zeile 5, Kopfzeile in HVM-Orange, freeze_panes A6,
Betraege mit zwei Nachkommastellen, Datumszellen als Datum, PLZ als Text, leere Zellen ohne Platzhalter."""

from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from apps.lists.data import AMOUNT_COLUMNS, DATE_COLUMNS, OPEN_COLUMNS, TEXT_COLUMNS, WRAP_COLUMNS, ListData

ORANGE = "E6A83C"
TEXT = "1A1A1A"
FONT = "Arial"
HEADER_ROW = 5
MAX_WIDTH = 60
EMPTY_MARKER = {
    "Offene Punkte": "Keine offenen Punkte",
    "Historie": "Keine Datensätze",
    "Aktuell": "Keine Datensätze",
}


def _table_name(title: str) -> str:
    return re.sub(r"[^A-Za-z0-9_]", "", title) or "Tabelle"


def _write_sheet(ws, title: str, columns: list[str], rows: list[dict], header: dict) -> None:
    ws.title = title
    ws["A1"] = f"Objekt {header['object_number']} {header['object_name']}".strip()
    ws["A1"].font = Font(name=FONT, bold=True, size=12)
    ws["A2"] = header["management_type"]
    ws["A3"] = f"Stand: {header['generated_at'].strftime('%d.%m.%Y %H:%M')}, {len(rows)} Datensätze"
    for cell in ("A2", "A3"):
        ws[cell].font = Font(name=FONT, size=10)
    for col, name in enumerate(columns, start=1):
        c = ws.cell(row=HEADER_ROW, column=col, value=name)
        c.font = Font(name=FONT, bold=True, color=TEXT, size=10)
        c.fill = PatternFill("solid", fgColor=ORANGE)
        c.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {name: len(name) for name in columns}
    data_rows = rows or [{columns[0]: EMPTY_MARKER.get(title, "Keine Datensätze")}]
    for r_idx, row in enumerate(data_rows, start=HEADER_ROW + 1):
        for col, name in enumerate(columns, start=1):
            value = row.get(name)
            if value in (None, ""):
                continue
            c = ws.cell(row=r_idx, column=col)
            c.font = Font(name=FONT, size=10)
            if name in AMOUNT_COLUMNS and isinstance(value, Decimal | int | float):
                c.value = float(value)
                c.number_format = "#,##0.00"
                text_len = len(f"{value:,.2f}")
            elif name in DATE_COLUMNS and isinstance(value, date | datetime):
                c.value = (
                    value if isinstance(value, datetime) else datetime(value.year, value.month, value.day)
                )
                c.number_format = "DD.MM.YYYY"
                text_len = 10
            elif name in TEXT_COLUMNS:
                c.value = str(value)
                c.number_format = "@"
                text_len = len(str(value))
            elif isinstance(value, Decimal):
                c.value = float(value)
                c.number_format = "0.0000"
                text_len = len(str(value))
            else:
                c.value = value if isinstance(value, int | float) else str(value)
                text_len = len(str(value))
            if name in WRAP_COLUMNS:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            widths[name] = max(widths[name], min(text_len, MAX_WIDTH))
    for col, name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col)].width = min(widths[name] + 2, MAX_WIDTH)
    last_row = HEADER_ROW + len(data_rows)
    ref = f"A{HEADER_ROW}:{get_column_letter(len(columns))}{last_row}"
    table = Table(displayName=_table_name(title), ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleLight1", showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)
    ws.freeze_panes = f"A{HEADER_ROW + 1}"


def write_xlsx(data: ListData, path: Path, *, include_provenance: bool = False) -> Path:
    wb = Workbook()
    wb.properties.title = f"{data.title.replace('ü', 'ue')} {data.header['object_number']}"
    wb.properties.creator = "Hausverwaltung Müller GmbH, Objektakte"
    wb.properties.created = data.header["generated_at"].replace(tzinfo=None)
    ws = wb.active
    _write_sheet(ws, "Aktuell", data.columns, data.current, data.header)
    _write_sheet(wb.create_sheet(), "Historie", data.columns, data.history, data.header)
    _write_sheet(wb.create_sheet(), "Offene Punkte", OPEN_COLUMNS, data.open_items, data.header)
    if include_provenance and data.provenance:
        ws_p = wb.create_sheet()
        _write_sheet(
            ws_p,
            "Herkunft",
            ["Einheit", "Eigentümer", "Feld", "Quelle", "Status"],
            data.provenance,
            data.header,
        )
        ws_p.sheet_state = "hidden"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path
