"""Formatprofile des Imports (Fachentwurf H 6.2; M3 Teil 1: csv_generic, xlsx_generic, immoware24_export, generic_table).

Jedes Profil liefert einen Erkennungsscore und eine RawTable (alle Zeilen als Zeichenketten, Kopfzeile wird in
apps.imports.mapping bestimmt). Leerzeilen und Summenzeilen bleiben erhalten und werden spaeter als not_a_record
gefuehrt (nichts wird verworfen). PDF-Profile folgen nach der OCR-Pipeline (M9).
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
DELIMITERS = ";,\t|"
IMMOWARE24_EINHEITEN = [
    "Objekt-Nr",
    "Status",
    "Objekt",
    "Verwaltungsart",
    "Gebaeude",
    "VE-Nr",
    "VE-Beschreibung",
    "Lage",
    "Eigentuemer",
    "Hausgeld_EUR_mtl",
    "Mieter",
    "Miete_EUR_mtl",
]
IMMOWARE24_KONTAKTE = ["ID", "Name", "Briefanrede", "Adresse", "PLZ", "Stadt", "Land", "Telefon", "E-Mail"]
UNIT_LIKE = re.compile(
    r"^\s*(WE|GE|TG|GA|ST|SP|S|Garage|Stellplatz|Wohnung|Haus)\s*(Nr\.?)?\s*\d+", re.IGNORECASE
)


@dataclass
class RawTable:
    rows: list[list[str]]
    sheet: str | None = None
    meta: dict = field(default_factory=dict)

    @property
    def width(self) -> int:
        return max((len(r) for r in self.rows), default=0)


def _cell_to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "ja" if value else "nein"
    if isinstance(value, datetime):
        return (
            value.strftime("%d.%m.%Y")
            if value.time() == datetime.min.time()
            else value.strftime("%d.%m.%Y %H:%M")
        )
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{Decimal(str(value)).normalize():f}".replace(".", ",")
    if isinstance(value, Decimal):
        return f"{value.normalize():f}".replace(".", ",")
    return str(value).strip()


def detect_encoding(data: bytes) -> str:
    for enc in ENCODINGS:
        try:
            data.decode(enc)
            return enc
        except UnicodeDecodeError:
            continue
    return "latin-1"


def detect_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=DELIMITERS).delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in DELIMITERS}
        return max(counts, key=counts.get) if any(counts.values()) else ";"


def read_csv(path: Path) -> RawTable:
    data = Path(path).read_bytes()
    enc = detect_encoding(data)
    text = data.decode(enc)
    delimiter = detect_delimiter("\n".join(text.splitlines()[:20]))
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [[c.strip() for c in row] for row in reader]
    return RawTable(rows, None, {"encoding": enc, "delimiter": delimiter})


def _sheet_score(rows: list[list[str]]) -> int:
    return sum(1 for r in rows if any(UNIT_LIKE.match(c or "") for c in r))


def read_xlsx(path: Path, sheet: str | None = None) -> RawTable:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    tables: dict[str, list[list[str]]] = {}
    for ws in wb.worksheets:
        values = [[_cell_to_text(c) for c in row] for row in ws.iter_rows(values_only=True)]
        # verbundene Zellen: Wert der linken oberen Zelle auf alle Zellen des Bereichs uebertragen
        for rng in ws.merged_cells.ranges:
            top = values[rng.min_row - 1][rng.min_col - 1] if rng.min_row - 1 < len(values) else ""
            for r in range(rng.min_row - 1, min(rng.max_row, len(values))):
                for c in range(rng.min_col - 1, rng.max_col):
                    if c < len(values[r]):
                        values[r][c] = top
        tables[ws.title] = values
    if not tables:
        return RawTable([], None, {"sheets": []})
    chosen = (
        sheet if sheet in tables else max(tables, key=lambda t: (_sheet_score(tables[t]), len(tables[t])))
    )
    return RawTable(
        tables[chosen], chosen, {"sheets": list(tables), "sheet_rows": {t: len(v) for t, v in tables.items()}}
    )


class FormatProfile:
    code = ""
    source_format = ""
    extensions: tuple[str, ...] = ()

    def detect(self, path: Path, filename: str) -> float:
        raise NotImplementedError

    def extract(self, path: Path, **options) -> RawTable:
        raise NotImplementedError


class CsvGeneric(FormatProfile):
    code = "csv_generic"
    source_format = "csv"
    extensions = (".csv", ".txt", ".tsv")

    def detect(self, path, filename):
        return 0.6 if filename.lower().endswith(self.extensions) else 0.0

    def extract(self, path, **options):
        return read_csv(path)


class XlsxGeneric(FormatProfile):
    code = "xlsx_generic"
    source_format = "xlsx"
    extensions = (".xlsx", ".xlsm")

    def detect(self, path, filename):
        return 0.7 if filename.lower().endswith(self.extensions) else 0.0

    def extract(self, path, **options):
        return read_xlsx(path, options.get("sheet"))


def _header_hits(row: list[str], expected: list[str]) -> int:
    norm = {c.strip().lower() for c in row}
    return sum(1 for e in expected if e.lower() in norm)


class Immoware24Export(FormatProfile):
    """Export einheiten.csv (und kontakte.csv) mit den Spalten aus Befund 4 (H 6.2.1)."""

    code = "immoware24_export"
    source_format = "immoware24"
    extensions = (".csv", ".txt")

    def detect(self, path, filename):
        if not filename.lower().endswith(self.extensions):
            return 0.0
        try:
            table = read_csv(path)
        except Exception:
            return 0.0
        for row in table.rows[:5]:
            if _header_hits(row, IMMOWARE24_EINHEITEN) >= 8 or _header_hits(row, IMMOWARE24_KONTAKTE) >= 7:
                return 0.95
        return 0.0

    def extract(self, path, **options):
        table = read_csv(path)
        for row in table.rows[:5]:
            if _header_hits(row, IMMOWARE24_KONTAKTE) >= 7:
                table.meta["immoware24_file"] = "kontakte"
                break
            if _header_hits(row, IMMOWARE24_EINHEITEN) >= 8:
                table.meta["immoware24_file"] = "einheiten"
                break
        return table


class GenericTable(FormatProfile):
    code = "generic_table"
    source_format = "other"

    def detect(self, path, filename):
        return 0.1

    def extract(self, path, **options):
        name = str(path).lower()
        if name.endswith((".xlsx", ".xlsm")):
            return read_xlsx(path, options.get("sheet"))
        return read_csv(path)


PROFILES: dict[str, FormatProfile] = {
    p.code: p for p in (Immoware24Export(), XlsxGeneric(), CsvGeneric(), GenericTable())
}


def choose_profile(path: Path, filename: str) -> tuple[FormatProfile, dict[str, float]]:
    scores = {code: p.detect(Path(path), filename) for code, p in PROFILES.items()}
    best = max(scores, key=lambda c: (scores[c], c == "immoware24_export"))
    return PROFILES[best], scores
