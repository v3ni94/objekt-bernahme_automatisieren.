"""Importprotokoll als Excel (H 6.8): Blatt Zeilen und Blatt Zusammenfassung unter /srv/objektakte/imports/<batch_id>/.
Enthaelt keine vollstaendige IBAN (raw_data ist maskiert, parsed_fields tragen nur iban_last4 und Hash)."""

from __future__ import annotations

import json
from pathlib import Path

from apps.imports.models import ImportBatch, ImportRow
from apps.imports.services import import_dir


def write_protocol(batch: ImportBatch) -> Path:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Zeilen"
    header = [
        "Zeile",
        "Teil",
        "Status",
        "Konfidenz",
        "Gründe",
        "Zielart",
        "Person",
        "Einheit",
        "Vorschlag",
        "Entscheidung",
        "Bearbeiter",
        "Zeitpunkt",
        "Zielentitäten",
        "Bemerkung",
        "Rohzeile",
    ]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    rows = (
        ImportRow.objects.filter(batch=batch).select_related("confirmed_by").order_by("row_no", "sub_index")
    )
    for r in rows:
        pf = r.parsed_fields or {}
        person = pf.get("person") or {}
        shared = pf.get("shared") or {}
        ws.append(
            [
                r.row_no,
                r.sub_index,
                r.status,
                float(r.confidence) if r.confidence is not None else None,
                ", ".join(r.uncertainty_reasons or []),
                r.target_entity_type or "",
                person.get("company_name")
                or " ".join(p for p in [person.get("first_name"), person.get("last_name")] if p),
                (shared.get("unit_label") or {}).get("value") or "",
                pf.get("proposal") or "",
                json.dumps(
                    (r.committed_targets or [{}])[0].get("type") if r.committed_targets else "",
                    ensure_ascii=False,
                )
                if r.committed_targets
                else "",
                r.confirmed_by.email if r.confirmed_by else "",
                r.confirmed_at.strftime("%d.%m.%Y %H:%M") if r.confirmed_at else "",
                json.dumps(r.committed_targets or [], ensure_ascii=False),
                r.notes or "",
                json.dumps(r.raw_data, ensure_ascii=False),
            ]
        )
    ws2 = wb.create_sheet("Zusammenfassung")
    mapping = batch.column_mapping or {}
    for k, v in [
        ("Import", batch.pk),
        ("Objekt", batch.object.object_number),
        ("Datei", batch.source_file_name),
        ("SHA-256", batch.source_sha256),
        ("Profil", batch.parser_profile),
        ("Parser-Version", batch.parser_version),
        ("Status", batch.status),
        ("Zeilen gesamt", batch.rows_total),
        ("unsicher", batch.rows_uncertain),
        ("bestätigt", batch.rows_confirmed),
        ("abgelehnt", batch.rows_rejected),
        ("übernommen", batch.rows_committed),
        ("Blatt", mapping.get("sheet") or ""),
    ]:
        ws2.append([k, v])
    ws2.append([])
    ws2.append(["Spaltenzuordnung", "Zielfeld", "Konfidenz", "manuell"])
    for col in mapping.get("columns", []):
        ws2.append(
            [
                col.get("source_header"),
                col.get("target") or "nicht übernommen",
                col.get("confidence"),
                "ja" if col.get("manual") else "",
            ]
        )
    target_dir = import_dir(batch)
    target_dir.mkdir(parents=True, exist_ok=True)
    path = target_dir / "protokoll.xlsx"
    wb.save(path)
    return path
