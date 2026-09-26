"""Formatweiche und Seitenanalyse (Fachentwurf E 1.2 Schritte 4 bis 6, 1.3; docs/architektur.md 6.1 Nr. 3, 4).

Eine Seite gilt als digital, wenn die Textebene mindestens ocr.digital_min_chars Zeichen hat, der Anteil
alphanumerischer Zeichen ueber ocr.digital_alnum_ratio liegt und die Seite nicht nur aus einem seitenfuellenden Bild
mit wenig Text besteht (ANNAHME A-11). Rein digitale Dateien starten ocrmypdf nicht. Bilder werden verlustfrei mit
img2pdf in PDF gewandelt, Office-Dateien ohne OCR extrahiert, andere Formate gehen in die manuelle Pruefung.

26.09.2026: Endungslose Dateien und Temporaer-Endungen (.tmp, .herunterladen, .indir, .hed, .part, .crdownload)
werden am Inhalt erkannt (apps.pipeline.sniff), bevor sie als nicht unterstuetzt gelten; HTML laeuft als Textseite
ohne OCR. Office-Altformate (.doc, .xls, .rtf, .odt, .ods, Art office_legacy) wandelt LibreOffice im
Arbeitsverzeichnis in eine PDF (apps.pipeline.office_convert), die dann wie eine PDF analysiert wird (Textebene,
OCR nur wenn noetig, Seitenbilder); abgelegt bleibt das Original. Ohne LibreOffice oder bei gescheiterter
Umwandlung entsteht der Fall "nicht unterstuetztes Format" mit Notiz und content_checked; die vor dem Scheitern
erkannte Art steht in detected_kind, damit formate_wiederaufnehmen (Gruppe office) auch Faelle findet, deren
Altformat nur die Inhaltspruefung erkannt hat. Praesentationen (.ppt, .odp) bleiben unsupported: das Worker-Image
bringt kein libreoffice-impress mit, fuer Paperless-Dokumente greift weiter die PDF-Archivfassung.
"""

from __future__ import annotations

import csv
import logging
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path

from apps.pipeline import email_text, office_convert, sniff

PDF_MIMES = {"application/pdf"}
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".gif", ".webp"}
OFFICE_EXT = {".docx", ".xlsx", ".xlsm", ".csv", ".txt"}
EMAIL_EXT = email_text.EMAIL_EXT  # .eml, .msg: Textseite ohne OCR (25.09.2026)
EMAIL_MIMES = email_text.EMAIL_MIMES
HTML_EXT = {".html", ".htm"}  # Textseite ueber email_text.html_to_text (26.09.2026)
HTML_MIMES = {"text/html", "application/xhtml+xml"}
# Office-Altformate (26.09.2026): Umwandlung in PDF ueber LibreOffice (writer, calc), siehe office_convert.
# Kein .ppt und kein application/vnd.ms-powerpoint: das Worker-Image hat kein libreoffice-impress, soffice endete
# dann mit Rueckgabecode 0 ohne PDF; Praesentationen bleiben unsupported (Paperless-Archivfassung als Ersatz)
OFFICE_LEGACY_EXT = {".doc", ".xls", ".rtf", ".odt", ".ods"}
OFFICE_LEGACY_MIMES = {
    "application/msword",
    "application/vnd.ms-excel",
    "application/rtf",
    "text/rtf",
    "application/vnd.oasis.opendocument.text",
    "application/vnd.oasis.opendocument.spreadsheet",
}
TEMP_SUFFIXES = sniff.TEMP_SUFFIXES
TEXT_KINDS = ("office", "email", "html")  # digitaler Text ohne OCR, eine Seite je Zeichenkette
# Arten, die die Kette nicht liest: Fall "nicht unterstuetztes Format", Paperless-Archivfassung als Ersatz.
# 26.09.2026: office_legacy entfaellt hier, die Kette wandelt Altformate selbst (Umwandlung scheitert: unsupported)
UNREADABLE_KINDS = ("unsupported",)
# Arten, deren Seiten aus einer (gewandelten) PDF stammen: origin_kind aus Textebene und OCR-Seiten
PDF_BASED_KINDS = ("pdf", "image", "office_legacy")
GOOGLE_DOC_PREFIX = "application/vnd.google-apps."
logger = logging.getLogger(__name__)


@dataclass
class PageInfo:
    page_no: int
    chars: int
    alnum_ratio: float
    image_cover: float
    digital: bool
    reason: str


@dataclass
class Analysis:
    kind: str  # pdf | image | office | email | html | office_legacy (gewandelt) | google_doc | unsupported
    page_count: int = 0
    pages: list[PageInfo] = field(default_factory=list)
    digital_pages: list[int] = field(default_factory=list)
    ocr_pages: list[int] = field(default_factory=list)
    chunks: list[list[int]] = field(default_factory=list)
    texts: dict[int, str] = field(
        default_factory=dict
    )  # Textebene digitaler Seiten (unmaskiert, nur im Speicher)
    pdf_path: str | None = None
    note: str | None = None
    content_checked: bool = (
        False  # Inhalt geprueft (sniff), Kennzeichen fuer die spaetere Bereinigung (26.09.2026)
    )
    suffix: str | None = None  # massgebliche Endung (aus dem Inhalt, sonst die der Datei), 26.09.2026
    # Art aus der Formatweiche vor einer gescheiterten Umwandlung (office_legacy bei kind unsupported mit Notiz
    # "Umwandlung ..."); formate_wiederaufnehmen (Gruppe office) nimmt solche Faelle darueber erneut auf, auch
    # wenn detect_kind nach Name und MIME-Typ nichts erkennt (endungslose Datei mit RTF- oder OLE-Inhalt), 26.09.2026
    detected_kind: str | None = None
    # Ergebnis der Inhaltspruefung getrennt von content_checked (26.09.2026): None = Inhalt nicht geprueft,
    # False = keine Signatur erkannt (Temporaerdatei, Kandidat fuer restformate_bereinigen), True = Signatur erkannt,
    # auch wenn die Art nicht verarbeitbar ist (Praesentation, ZIP-Archiv, Office-Paket ohne bekannten Hauptteil).
    # content_suffix ist die aus dem Inhalt abgeleitete Endung (".ppt", ".zip", "" ohne Vorschlag), sonst None.
    content_recognized: bool | None = None
    content_suffix: str | None = None

    @property
    def origin_kind(self) -> str | None:
        if self.kind not in PDF_BASED_KINDS:
            return "digital" if self.kind in (*TEXT_KINDS, "google_doc") else None
        if self.ocr_pages and self.digital_pages:
            return "mixed"
        return "scan" if self.ocr_pages else "digital"

    def to_json(self) -> dict:
        d = asdict(self)
        d.pop("texts", None)
        return d


def detect_kind(path: Path, mime_type: str | None) -> str:
    ext = path.suffix.lower()
    mime = (mime_type or "").lower()
    if mime in PDF_MIMES or ext == ".pdf":
        return "pdf"
    if mime.startswith(GOOGLE_DOC_PREFIX):
        return "google_doc"
    if ext in IMAGE_EXT or mime.startswith("image/"):
        return "image"
    if ext in OFFICE_EXT:
        return "office"
    if ext in EMAIL_EXT or mime in EMAIL_MIMES:
        return "email"
    if ext in HTML_EXT or mime in HTML_MIMES:
        return "html"
    if ext in OFFICE_LEGACY_EXT or mime in OFFICE_LEGACY_MIMES:
        return "office_legacy"
    return "unsupported"


def detect_kind_with_content(
    path: Path, mime_type: str | None
) -> tuple[str, str, bool, str | None, sniff.Sniffed | None]:
    """Formatweiche mit Inhaltspruefung (26.09.2026): Rueckgabe Art, massgebliche Endung, ob der Inhalt geprueft
    wurde, eine Notiz und das Ergebnis der Inhaltspruefung (None, wenn nicht geprueft oder keine Signatur erkannt).
    Bei unbekannter oder temporaerer Endung zaehlt nur der Inhalt, nicht der MIME-Typ."""
    kind = detect_kind(path, mime_type)
    suffix = path.suffix.lower()
    if kind == "google_doc":
        return kind, suffix, False, None, None
    if suffix not in TEMP_SUFFIXES and kind != "unsupported":
        return kind, suffix, False, None, None
    sniffed = sniff.detect_kind_by_content(path)
    if sniffed is None:
        return (
            "unsupported",
            suffix,
            True,
            "Format nicht unterstützt, manuelle Prüfung (Inhalt unbekannt)",
            None,
        )
    if sniffed.kind == "unsupported":
        label = sniffed.suffix.lstrip(".") or "unbekannt"
        return (
            "unsupported",
            suffix,
            True,
            f"Format nicht unterstützt, manuelle Prüfung (Inhalt: {label})",
            sniffed,
        )
    note = f"Format aus Dateiinhalt erkannt: {sniffed.kind}"
    if sniffed.suffix:
        note += f" ({sniffed.suffix})"
    return sniffed.kind, sniffed.suffix or suffix, True, note, sniffed


def image_to_pdf(image_path: Path, target: Path) -> Path:
    import img2pdf
    from PIL import Image

    with Image.open(image_path) as im:
        frames = getattr(im, "n_frames", 1)
        if im.mode in ("RGBA", "P", "LA") or frames > 1:
            pages = []
            for i in range(frames):
                im.seek(i)
                buf = target.with_suffix(f".p{i}.png")
                im.convert("RGB").save(buf)
                pages.append(buf)
            target.write_bytes(img2pdf.convert([str(p) for p in pages], rotation=img2pdf.Rotation.ifvalid))
            for p in pages:
                p.unlink(missing_ok=True)
            return target
    # rotation=ifvalid (23.09.2026): 24 Scans trugen EXIF-Drehung 0 und scheiterten mit ExifOrientationError
    target.write_bytes(img2pdf.convert(str(image_path), rotation=img2pdf.Rotation.ifvalid))
    return target


def _alnum_ratio(text: str) -> float:
    stripped = re.sub(r"\s", "", text)
    if not stripped:
        return 0.0
    return sum(ch.isalnum() for ch in stripped) / len(stripped)


def _image_cover(pdf, page_index: int) -> float:
    """Anteil der Seitenflaeche, den das groesste Bildobjekt bedeckt (Bounding Box aus dem Inhaltsstrom)."""
    try:
        page = pdf[page_index]
        width, height = page.get_width(), page.get_height()
        if not width or not height:
            return 0.0
        best = 0.0
        for obj in page.get_objects(filter=(3,)):  # 3 = FPDF_PAGEOBJ_IMAGE
            getter = getattr(obj, "get_bounds", None) or getattr(obj, "get_pos", None)
            if getter is None:
                continue
            left, bottom, right, top = getter()
            area = max(0.0, right - left) * max(0.0, top - bottom)
            best = max(best, area / (width * height))
        return min(best, 1.0)
    except Exception:
        logger.debug("Bildflaeche der Seite %s nicht bestimmbar", page_index + 1, exc_info=True)
        return 0.0


def analyze_pdf(
    path: Path, *, min_chars: int, alnum_ratio: float, cover_ratio: float, chunk_pages: int
) -> Analysis:
    import pypdfium2 as pdfium

    result = Analysis(kind="pdf", pdf_path=str(path))
    pdf = pdfium.PdfDocument(str(path))
    try:
        result.page_count = len(pdf)
        for i in range(result.page_count):
            page = pdf[i]
            textpage = page.get_textpage()
            text = textpage.get_text_bounded() or ""
            textpage.close()
            chars = len(text.strip())
            ratio = _alnum_ratio(text)
            cover = _image_cover(pdf, i)
            digital, reason = True, "text_layer"
            if chars < min_chars:
                digital, reason = False, "too_few_chars"
            elif ratio < alnum_ratio:
                digital, reason = False, "garbage_text_layer"
            elif cover >= cover_ratio and chars < max(min_chars * 4, 200):
                digital, reason = False, "full_page_image"
            info = PageInfo(i + 1, chars, round(ratio, 3), round(cover, 3), digital, reason)
            result.pages.append(info)
            if digital:
                result.digital_pages.append(i + 1)
                result.texts[i + 1] = text
            else:
                result.ocr_pages.append(i + 1)
            page.close()
    finally:
        pdf.close()
    result.chunks = [
        result.ocr_pages[i : i + chunk_pages] for i in range(0, len(result.ocr_pages), chunk_pages)
    ]
    return result


def extract_office_text(path: Path, suffix: str | None = None) -> list[str]:
    """Textextraktion ohne OCR: docx je Absatz, xlsx je Blatt, csv und txt als eine Seite. suffix uebersteuert die
    Endung der Datei, wenn die Art aus dem Inhalt erkannt wurde (26.09.2026: lokale Kopie heisst dann "original")."""
    ext = (suffix or path.suffix).lower()
    if ext == ".docx":
        import docx

        document = docx.Document(str(path))
        text = "\n".join(p.text for p in document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                text += "\n" + "\t".join(c.text for c in row.cells)
        return [text]
    if ext in (".xlsx", ".xlsm"):
        import openpyxl

        # Dateiobjekt statt Pfad: openpyxl prueft sonst die Endung und lehnt "original" ohne .xlsx ab (26.09.2026)
        with path.open("rb") as fh:
            wb = openpyxl.load_workbook(fh, read_only=True, data_only=True)
            pages = []
            for ws in wb.worksheets:
                lines = [
                    "\t".join("" if c is None else str(c) for c in row)
                    for row in ws.iter_rows(values_only=True)
                ]
                pages.append(f"[Blatt {ws.title}]\n" + "\n".join(lines))
            wb.close()
        return pages or [""]
    if ext == ".csv":
        data = path.read_bytes()
        for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(text.splitlines(), dialect))
        return ["\n".join("\t".join(r) for r in rows)]
    return [path.read_text(encoding="utf-8", errors="replace")]


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_html_text(path: Path) -> list[str]:
    """HTML als eine Textseite (26.09.2026): Zeichensatz aus dem Dokument, sonst Reihe utf-8, cp1252, latin-1;
    Skripte, Stile und Kopf entfallen (email_text.html_to_text), Kuerzung wie bei E-Mail-Koerpern."""
    data = path.read_bytes()
    text = None
    match = re.search(rb"charset=[\"']?([A-Za-z0-9_-]+)", data[:4096])
    if match:
        try:
            text = data.decode(match.group(1).decode("ascii"), errors="strict")
        except (LookupError, UnicodeDecodeError):
            text = None
    if text is None:
        text = _decode_text(data)
    text = email_text.html_to_text(text)
    if len(text) > email_text.MAX_TEXT_CHARS:
        text = text[: email_text.MAX_TEXT_CHARS] + "\n[Text gekürzt]"
    return [text]


def convert_legacy_office(path: Path, work: Path, *, timeout_s: int) -> Path:
    """Office-Altformat in work/converted.pdf wandeln (26.09.2026). LibreOffice schreibt <stem>.pdf in das
    Ausgabeverzeichnis; ein Unterverzeichnis verhindert, dass original.pdf neben original.doc liegt und
    storage.original_path die falsche Datei findet. Ausnahmen aus office_convert laufen durch."""
    out_dir = work / "convert"
    produced = office_convert.convert_to_pdf(path, out_dir, timeout_s)
    target = work / "converted.pdf"
    target.unlink(missing_ok=True)
    produced.replace(target)
    return target


def _extract_text_pages(kind: str, path: Path, suffix: str) -> list[str]:
    if kind == "office":
        return extract_office_text(path, suffix)
    if kind == "email":
        return email_text.extract_email_text(path, suffix)
    return extract_html_text(path)


def analyze_file(
    path: Path,
    mime_type: str | None,
    *,
    min_chars: int,
    alnum_ratio: float,
    cover_ratio: float,
    chunk_pages: int,
    work: Path,
    office_timeout_s: int = office_convert.DEFAULT_TIMEOUT_S,
) -> Analysis:
    kind, suffix, content_checked, note, sniffed = detect_kind_with_content(path, mime_type)
    # Ergebnis der Inhaltspruefung fuer restformate_bereinigen festhalten (26.09.2026): nur Dateien ohne erkannte
    # Signatur sind Temporaerdateien; content_checked allein reicht nicht, weil es auch nach einer gescheiterten
    # Office-Umwandlung gesetzt wird
    content_recognized = (sniffed is not None) if content_checked else None
    content_suffix = sniffed.suffix if sniffed is not None else None
    if kind == "pdf":
        result = analyze_pdf(
            path,
            min_chars=min_chars,
            alnum_ratio=alnum_ratio,
            cover_ratio=cover_ratio,
            chunk_pages=chunk_pages,
        )
    elif kind == "image":
        pdf_path = image_to_pdf(path, work / "converted.pdf")
        result = analyze_pdf(
            pdf_path,
            min_chars=min_chars,
            alnum_ratio=alnum_ratio,
            cover_ratio=cover_ratio,
            chunk_pages=chunk_pages,
        )
        result.kind = "image"
    elif kind in TEXT_KINDS:
        texts = _extract_text_pages(kind, path, suffix)
        result = Analysis(kind=kind, page_count=len(texts), digital_pages=list(range(1, len(texts) + 1)))
        result.texts = {i + 1: t for i, t in enumerate(texts)}
        result.pages = [
            PageInfo(i + 1, len(t), round(_alnum_ratio(t), 3), 0.0, True, f"{kind}_text")
            for i, t in enumerate(texts)
        ]
    elif kind == "office_legacy":
        # Umwandlung ueber LibreOffice (26.09.2026); danach Analyse wie PDF. Scheitert sie, Fall "nicht
        # unterstuetztes Format" mit Notiz; content_checked, damit formate_wiederaufnehmen (Gruppe office) den
        # Fall gezielt und nicht ueber die Gruppen bekannt oder temporaer erneut aufnimmt
        try:
            pdf_path = convert_legacy_office(path, work, timeout_s=office_timeout_s)
        except office_convert.ConversionUnavailable as exc:
            logger.warning("Office-Altformat nicht gewandelt: %s", exc)
            result = Analysis(kind="unsupported", note="Umwandlung nicht verfügbar")
            content_checked = True
        except office_convert.ConversionFailed as exc:
            logger.warning("Office-Altformat nicht gewandelt: %s", exc)
            result = Analysis(kind="unsupported", note=f"Umwandlung fehlgeschlagen ({exc})")
            content_checked = True
        else:
            result = analyze_pdf(
                pdf_path,
                min_chars=min_chars,
                alnum_ratio=alnum_ratio,
                cover_ratio=cover_ratio,
                chunk_pages=chunk_pages,
            )
            result.kind = "office_legacy"
    elif kind == "google_doc":
        result = Analysis(kind=kind, note="Google-Dokument, Export folgt in M6")
    else:
        result = Analysis(kind="unsupported", note=note or "Format nicht unterstützt, manuelle Prüfung")
    result.content_checked = content_checked
    result.content_recognized = content_recognized
    result.content_suffix = content_suffix
    result.suffix = suffix or None
    result.detected_kind = kind
    if note and result.note is None:
        result.note = note
    return result
