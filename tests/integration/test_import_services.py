"""Importkette Ende zu Ende ohne Oberflaeche: Immoware24-Profil, generische CSV und XLSX, Konflikt mit bestehender
Zuordnung, IBAN-Maskierung, rows_total gleich Summe, 1.000 Zeilen unter zwei Minuten (ANNAHME A-21)."""

from __future__ import annotations

import json
import time
from datetime import date
from pathlib import Path

import pytest

from apps.imports import services
from apps.imports.models import ImportBatch, ImportRow
from apps.imports.protocol import write_protocol
from apps.imports.services import Decision
from apps.objects.models import ManagedObject, Unit
from apps.parties.models import FieldProvenance, Owner, OwnerUnitAssignment, Tenant
from apps.review.models import ReviewCase

pytestmark = pytest.mark.django_db
FIX = Path(__file__).resolve().parents[1] / "fixtures" / "imports"
FULL_IBAN = "DE02120300000000202051"


@pytest.fixture
def objekt(seeded):
    return ManagedObject.objects.create(
        object_number="623", name="Musterstadt, Musterweg 1", management_type="weg", is_test=True
    )


def run_chain(objekt, filename: str, data: bytes | None = None, user=None) -> ImportBatch:
    payload = data if data is not None else (FIX / filename).read_bytes()
    batch, created = services.create_batch(
        objekt, filename=filename, data=payload, user=user, import_kind="mixed"
    )
    assert created
    services.parse_batch(batch)
    services.normalize_batch(batch, user=user)
    batch.refresh_from_db()
    return batch


def test_immoware24_kette(objekt, admin_user):
    batch = run_chain(objekt, "immoware24_einheiten.csv", user=admin_user)
    assert batch.parser_profile == "immoware24_export" and batch.status == "in_review"
    mapping = batch.column_mapping
    assert all(c["target"] for c in mapping["columns"]), "alle Spalten ohne manuelle Korrektur zugeordnet"
    rows = list(ImportRow.objects.filter(batch=batch).order_by("row_no", "sub_index"))
    assert services.status_sum_matches(batch) and batch.rows_total == len(rows)
    by_no = {}
    for r in rows:
        by_no.setdefault(r.row_no, []).append(r)
    # Zeile 2: Mustermann, Erika & Max -> zwei Personen, sicher
    assert [r.parsed_fields["person"]["first_name"] for r in by_no[2]] == ["Erika", "Max"] and all(
        r.status == "parsed" for r in by_no[2]
    )
    # Zeile 3: Leerzeichenform (mittel, kein Grund) plus Mieter mit Miete
    assert len(by_no[3]) == 3 and by_no[3][2].target_entity_type == "tenant_unit_assignment"
    # Zeile 5: Erbengemeinschaft, Lage aus Bezeichnung
    assert (
        by_no[5][0].parsed_fields["person"]["type"] == "community"
        and by_no[5][0].parsed_fields["shared"]["location"]["value"] == "1.OG rechts"
    )
    # Zeile 7: Status archiv und unsicherer Name
    assert by_no[7][0].status == "uncertain" and {"unit_status_not_active", "name_split_ambiguous"} <= set(
        by_no[7][0].uncertainty_reasons
    )
    assert by_no[7][0].parsed_fields["proposal"] == "reject"
    # Zeile 8: Titel und c/o
    p8 = by_no[8][0].parsed_fields
    assert (
        p8["person"]["last_name"] == "Neumuster-Beispiel"
        and p8["shared"]["correspondence_addition"]["value"] == "c/o Beispiel Hausverwaltung"
    )
    # Zeile 9: anderes Objekt
    assert (
        "other_object" in by_no[9][0].uncertainty_reasons
        and by_no[9][0].parsed_fields["proposal"] == "reject"
    )
    # Zeile 10: technisch, ohne Eigentuemer
    assert "unit_status_not_active" in by_no[10][0].uncertainty_reasons
    assert (
        ReviewCase.objects.filter(case_type="import_row_uncertain", batch_key=f"import:{batch.pk}").count()
        == batch.rows_uncertain
    )
    assert Owner.objects.count() == 0 and Unit.objects.count() == 0, (
        "Stammtabellen bleiben bis zur Bestätigung leer"
    )

    # Uebernahme: sichere Zeilen bestaetigen, unsichere mit Vorschlag reject ablehnen
    decisions = {}
    for r in rows:
        if r.status == "parsed":
            decisions[r.pk] = Decision("accept")
        elif r.status == "uncertain" and r.parsed_fields.get("proposal") == "reject":
            decisions[r.pk] = Decision("reject", reason="Vorschlag nicht übernehmen")
    result = services.commit_rows(batch, decisions, user=admin_user)
    batch.refresh_from_db()
    fehler = list(
        ImportRow.objects.filter(batch=batch, notes__startswith="Fehler").values_list(
            "row_no", "sub_index", "notes"
        )
    )
    assert result["failed"] == 0 and result["committed"] > 0, fehler
    assert Unit.objects.filter(object=objekt).count() == 6  # WE 1 bis WE 4, Garage 1, WE 6
    assert Unit.objects.get(object=objekt, external_ref="1004").unit_label == "WE 4"
    garage = Unit.objects.get(object=objekt, unit_label="Garage 1")
    assert garage.unit_type == "garage" and garage.external_ref == "1005" and garage.building == "Haus B"
    owners = {str(o): o for o in Owner.objects.all()}
    assert {
        "Erika Mustermann",
        "Max Mustermann",
        "Erika Beispiel",
        "Max Beispiel",
        "Muster GbR",
        "Erbengemeinschaft Altmuster",
        "Erika Neumuster-Beispiel",
    } <= set(owners)
    # Mustermann Erika hat WE 1 und Garage 1 mit einem Eigentuemerdatensatz (Zusammenfuehrung innerhalb des Imports)
    erika = owners["Erika Mustermann"]
    assert OwnerUnitAssignment.objects.filter(owner=erika).count() == 2
    assert (
        OwnerUnitAssignment.objects.filter(data_status="confirmed").count()
        == OwnerUnitAssignment.objects.count()
    )
    assert Tenant.objects.filter(last_name="Altmieter").exists()
    assert (
        FieldProvenance.objects.filter(
            entity_type="unit", field_name="unit_label", source_kind="import_row"
        ).count()
        == 6
    )
    assert batch.status == "committed" and services.status_sum_matches(batch)
    assert batch.rows_committed + batch.rows_rejected == batch.rows_total
    assert ReviewCase.objects.filter(batch_key=f"import:{batch.pk}", status="open").count() == 0


def test_generic_csv_mit_iban_und_summenzeile(objekt, admin_user):
    batch = run_chain(objekt, "eigentuemerliste_generic.csv", user=admin_user)
    rows = list(ImportRow.objects.filter(batch=batch).order_by("row_no", "sub_index"))
    total = [r for r in rows if r.uncertainty_reasons and "not_a_record" in r.uncertainty_reasons]
    assert len(total) == 1 and total[0].status == "rejected"
    dump = json.dumps([r.raw_data for r in rows] + [r.parsed_fields for r in rows], ensure_ascii=False)
    assert FULL_IBAN not in dump and "2020 51" not in dump and "202051" not in dump
    we01 = next(r for r in rows if r.row_no == 2)
    assert (
        we01.parsed_fields["shared"]["iban_last4"]["value"] == "2051"
        and we01.parsed_fields["shared"]["street"]["value"] == "Musterweg"
    )
    assert (
        we01.parsed_fields["shared"]["house_number"]["value"] == "12a"
        and we01.parsed_fields["shared"]["valid_from"]["value"] == "2015-01-01"
    )
    assert (
        we01.parsed_fields["shared"]["co_ownership_share"]["value"] == "125.00"
        and we01.parsed_fields["shared"]["co_ownership_share_base"]["value"] == 10000
    )
    services.commit_rows(
        batch, {r.pk: Decision("accept") for r in rows if r.status == "parsed"}, user=admin_user
    )
    o = Owner.objects.get(last_name="Mustermann")
    assert o.iban_last4 == "2051" and o.iban_hash is not None and o.correspondence_house_number == "12a"
    a = OwnerUnitAssignment.objects.get(owner=o)
    assert a.valid_from == date(2015, 1, 1) and a.unit.co_ownership_share_base == 10000
    path = write_protocol(batch)
    import openpyxl

    wb = openpyxl.load_workbook(path)
    text = " ".join(
        str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None
    )
    assert FULL_IBAN not in text and "2051" in text and "Zusammenfassung" in wb.sheetnames


def test_xlsx_getrennte_namensspalten(objekt, admin_user):
    batch = run_chain(objekt, "eigentuemerliste_generic.xlsx", user=admin_user)
    assert batch.parser_profile == "xlsx_generic" and batch.column_mapping["sheet"] == "Eigentümer"
    rows = list(ImportRow.objects.filter(batch=batch).exclude(status="rejected").order_by("row_no"))
    assert len(rows) == 3 and rows[0].parsed_fields["person"]["last_name"] == "Mustermann"
    services.commit_rows(batch, {r.pk: Decision("accept") for r in rows}, user=admin_user)
    hans = Owner.objects.get(last_name="Beispiel")
    assert OwnerUnitAssignment.objects.filter(owner=hans).count() == 2  # WE 2 und TG 1, ein Eigentuemer
    assert Unit.objects.get(object=objekt, unit_label="TG 1").unit_type == "underground_parking"
    assert Owner.objects.get(last_name="Mustermann").sepa_mandate_present is True


def test_konflikt_mit_bestehender_zuordnung_drei_optionen(objekt, admin_user):
    unit = Unit.objects.create(
        object=objekt, unit_label="WE 1", unit_label_normalized="WE1", unit_number="1", unit_type="apartment"
    )
    alt = Owner.objects.create(
        type="natural_person", last_name="Altmuster", first_name="Anna", search_name="ALTMUSTER ANNA"
    )
    from apps.parties.services import create_assignment

    create_assignment(owner=alt, unit=unit, valid_from=date(2010, 1, 1), valid_to=None, confirmed=True)
    csv = "Einheit;Eigentümer;Eigentumsbeginn\nWE 1;Neumuster, Nils;01.07.2026\nWE 1;Beispiel, Bernd;\nWE 1;Altmuster, Anna;\n".encode()
    batch = run_chain(objekt, "konflikt.csv", csv, user=admin_user)
    rows = list(ImportRow.objects.filter(batch=batch).order_by("row_no"))
    assert all(r.status == "uncertain" and "owner_conflict" in r.uncertainty_reasons for r in rows[:2])
    assert rows[0].parsed_fields["options"] == ["change_owner", "co_owner", "duplicate", "reject"]
    # Zeile 3 (Altmuster, Anna) trifft den bestehenden Eigentuemer automatisch (Vorname und Einheit)
    assert rows[2].parsed_fields["match"]["decision"] == "auto" and rows[2].status == "parsed"
    assert OwnerUnitAssignment.objects.count() == 1, "nichts überschrieben"

    services.commit_rows(batch, {rows[0].pk: Decision("change_owner")}, user=admin_user)
    alt_a = OwnerUnitAssignment.objects.get(owner=alt)
    assert alt_a.valid_to == date(2026, 6, 30)
    neu_a = OwnerUnitAssignment.objects.get(owner__last_name="Neumuster")
    assert neu_a.valid_from == date(2026, 7, 1) and neu_a.valid_to is None

    services.commit_rows(batch, {rows[1].pk: Decision("co_owner")}, user=admin_user)
    assert OwnerUnitAssignment.objects.filter(unit=unit, valid_to__isnull=True).count() == 2

    services.commit_rows(batch, {rows[2].pk: Decision("accept")}, user=admin_user)
    rows[2].refresh_from_db()
    assert rows[2].status == "committed" and Owner.objects.filter(last_name="Altmuster").count() == 1
    assert (
        OwnerUnitAssignment.objects.count() == 3
    )  # Altmuster (beendet), Neumuster, Beispiel; keine neue fuer Altmuster
    batch.refresh_from_db()
    assert batch.status == "committed" and services.status_sum_matches(batch)


def test_duplicate_option(objekt, admin_user):
    unit = Unit.objects.create(
        object=objekt, unit_label="WE 2", unit_label_normalized="WE2", unit_number="2", unit_type="apartment"
    )
    alt = Owner.objects.create(
        type="natural_person", last_name="Altmuster", first_name="Anna", search_name="ALTMUSTER ANNA"
    )
    from apps.parties.services import create_assignment

    create_assignment(owner=alt, unit=unit, valid_from=None, valid_to=None, confirmed=True)
    csv = "Einheit;Eigentümer\nWE 2;Neumuster, Nora\n".encode()
    batch = run_chain(objekt, "dublette.csv", csv, user=admin_user)
    row = ImportRow.objects.get(batch=batch)
    assert "owner_conflict" in row.uncertainty_reasons
    services.commit_rows(batch, {row.pk: Decision("duplicate", owner_id=alt.pk)}, user=admin_user)
    assert Owner.objects.count() == 1 and OwnerUnitAssignment.objects.count() == 1


def test_gleiche_datei_nur_einmal(objekt):
    data = (FIX / "eigentuemerliste_generic.csv").read_bytes()
    b1, c1 = services.create_batch(objekt, filename="a.csv", data=data)
    b2, c2 = services.create_batch(objekt, filename="a.csv", data=data)
    assert c1 and not c2 and b1.pk == b2.pk
    assert ImportBatch.objects.count() == 1


@pytest.mark.slow
def test_tausend_zeilen_unter_zwei_minuten(objekt, admin_user):
    lines = ["Einheit;Eigentümer;Straße;PLZ;Ort;Eigentumsbeginn"]
    for i in range(1, 1001):
        lines.append(f"WE {i};Mustermann{i}, Erika;Musterweg {i};12345;Musterstadt;01.01.2020")
    data = "\n".join(lines).encode()
    t0 = time.perf_counter()
    batch = run_chain(objekt, "gross.csv", data, user=admin_user)
    rows = ImportRow.objects.filter(batch=batch, status="parsed")
    assert rows.count() == 1000
    services.commit_rows(
        batch, {pk: Decision("accept") for pk in rows.values_list("pk", flat=True)}, user=admin_user
    )
    dauer = time.perf_counter() - t0
    assert Owner.objects.count() == 1000 and OwnerUnitAssignment.objects.count() == 1000
    assert dauer < 120, f"{dauer:.1f} s"
