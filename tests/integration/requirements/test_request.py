"""Nachforderung (M10, H 4): Entwurf ohne Frist traegt Wasserzeichen und kann nicht auf geprueft wechseln; Freigabe
erzeugt PDF ohne Wasserzeichen und DOCX mit identischem Positionsbestand, Audit mit Freigebendem; Textbaustein-
aenderung aendert bestehende Schreiben nicht; Anlage ab Schwelle; Versandvermerk; Oberflaeche."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pdfplumber
import pytest
from django.urls import reverse
from docx import Document as DocxDocument

from apps.audit.models import AuditEvent
from apps.config import store
from apps.requirements import engine
from apps.requirements import requests as rq
from apps.requirements.models import CompletenessFinding, DocumentRequest, RequestTextBlock

from .conftest import TODAY

pytestmark = pytest.mark.django_db


@pytest.fixture
def daten(objekt, szenario, settings, tmp_path):
    settings.OBJEKTAKTE = {
        **settings.OBJEKTAKTE,
        "DATA_DIR": tmp_path,
        "HVM_SIGNATURE_PATH": str(tmp_path / "keine.jpg"),
    }
    engine.evaluate_object(objekt, today=TODAY)
    return objekt


def pdf_text(path: str) -> str:
    with pdfplumber.open(path) as pdf:
        return "\n".join(p.extract_text() or "" for p in pdf.pages)


def docx_text(path: str) -> str:
    d = DocxDocument(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    for s in d.sections:
        parts.extend(p.text for p in s.header.paragraphs)
        parts.extend(p.text for p in s.footer.paragraphs)
    return "\n".join(parts)


def test_entwurf_pruefung_freigabe(daten, admin_user, clerk_user):
    obj = daten
    store.set("requests.attachment_threshold", 100, user=admin_user)  # Einzelaufstellung im Brieftext
    req = rq.create_request(obj, clerk_user)
    assert req.status == "draft" and req.version == 1 and req.deadline_date is None
    assert req.recipient_name == "Altverwaltung Beispiel GmbH" and req.reference == "VV-623"
    assert len(req.findings_snapshot) == len([i for i in engine.open_items(obj) if i["include_in_request"]])
    assert Path(req.pdf_path).exists() and Path(req.docx_path).exists() and req.content_hash
    text = pdf_text(req.pdf_path)
    assert rq.DRAFT_HINT in text and "[FRIST]" in text and "Nachforderung fehlender Unterlagen" in text
    with pdfplumber.open(req.pdf_path) as pdf:  # gedrehtes Wasserzeichen ENTWURF im Hintergrund
        assert "ENTWURF" in "".join(ch["text"] for ch in pdf.pages[0].chars)
    assert (
        "HRB 104762" in text
        and "Rheinpromenade 13" in text
        and "Einzelabrechnung 2025 für WE01 (Mustermann)" in text
    )
    assert "Timo Müller" in text and "Geschäftsführender Gesellschafter" in text
    assert "ENTWURF" in docx_text(req.docx_path)
    with pdfplumber.open(req.pdf_path) as pdf:
        assert pdf.pages[0].images, "Logo auf Seite 1"
    with pytest.raises(rq.RequestError):
        rq.mark_reviewed(req, clerk_user)
    with pytest.raises(rq.RequestError):
        rq.approve(req, admin_user)
    rq.update_request(req, clerk_user, deadline=date(2026, 8, 15))
    rq.mark_reviewed(req, clerk_user)
    req.refresh_from_db()
    assert (
        req.status == "reviewed" and req.reviewed_by == clerk_user and "15.08.2026" in pdf_text(req.pdf_path)
    )
    # Textbausteinaenderung nach der Erzeugung aendert das Schreiben nicht
    intro = RequestTextBlock.objects.get(code="intro")
    intro.text, intro.version = "geänderter Text", 2
    intro.save()
    body_before, hash_before = req.body, req.content_hash
    rq.approve(req, admin_user)
    req.refresh_from_db()
    assert req.status == "approved" and req.approved_by == admin_user and req.approved_at is not None
    assert req.body == body_before and req.content_hash == hash_before and "geänderter Text" not in req.body
    final = pdf_text(req.pdf_path)
    assert rq.DRAFT_HINT not in final and "[FRIST]" not in final
    assert req.approved_at.strftime("%d.%m.%Y") in final
    with pdfplumber.open(req.pdf_path) as pdf:
        assert "ENTWURF" not in "".join(ch["text"] for ch in pdf.pages[0].chars)
    assert "ENTWURF" not in docx_text(req.docx_path)
    positions = req.findings_snapshot
    dtext = docx_text(req.docx_path)
    assert all(p["line"] in dtext for p in positions) and all(
        p["line"] in final.replace("\n", " ") or p["unit_label"] in final for p in positions
    )
    assert AuditEvent.objects.filter(
        action="request.approve", entity_id=req.pk, user_id=admin_user.pk
    ).exists()
    with pytest.raises(rq.RequestError):
        rq.update_request(req, clerk_user, deadline=None)
    rq.mark_sent(req, clerk_user, note="per E-Mail an die Vorverwaltung")
    req.refresh_from_db()
    assert req.status == "sent" and req.sent_at is not None
    f = CompletenessFinding.objects.get(pk=positions[0]["finding_id"])
    assert f.details["requested_deadline"] == "2026-08-15" and f.details["request_version"] == 1
    assert any(i["requested_deadline"] == "2026-08-15" for i in engine.open_items(obj))
    rq.withdraw(req, admin_user, reason="Fehler in der Anschrift")
    req.refresh_from_db()
    assert req.status == "withdrawn"
    zweite = rq.create_request(obj, clerk_user, deadline=date(2026, 9, 1), reminder=True)
    assert zweite.version == 2 and "nehmen Bezug" in zweite.body and "01.09.2026" in pdf_text(zweite.pdf_path)


def test_anlage_ab_schwelle(daten, admin_user):
    store.set("requests.attachment_threshold", 2, user=admin_user)
    req = rq.create_request(daten, admin_user, deadline=date(2026, 8, 15))
    assert req.custom_text_blocks["attachment"] is True
    text = pdf_text(req.pdf_path)
    assert "Anlage: Einzelaufstellung der fehlenden Unterlagen" in text and "siehe Anlage" in text
    d = DocxDocument(req.docx_path)
    table = [t for t in d.tables if t.rows and t.rows[0].cells[0].text == "Einheit"][0]
    assert len(table.rows) == len(req.findings_snapshot) + 1


def test_ohne_offene_punkte_kein_schreiben(objekt, admin_user, settings, tmp_path):
    settings.OBJEKTAKTE = {**settings.OBJEKTAKTE, "DATA_DIR": tmp_path}
    with pytest.raises(rq.RequestError):
        rq.create_request(objekt, admin_user)


def test_oberflaeche(daten, admin_user, client_as):
    client = client_as(admin_user)
    resp = client.get(reverse("completeness", args=[daten.pk]))
    assert resp.status_code == 200 and "Offene Punkte" in resp.content.decode()
    resp = client.post(reverse("completeness", args=[daten.pk]), {"action": "evaluate"})
    assert resp.status_code == 302
    zv = CompletenessFinding.objects.get(object=daten, check_code="dunning_procedure", scope_type="object")
    resp = client.post(
        reverse("finding_override", args=[zv.pk]),
        {"manual_status": "fulfilled", "reason": "Negativerklärung liegt vor", "include_in_request": "1"},
    )
    assert resp.status_code == 302
    zv.refresh_from_db()
    assert zv.manual_status == "fulfilled"
    resp = client.get(reverse("request_list", args=[daten.pk]))
    assert resp.status_code == 200
    resp = client.post(reverse("request_list", args=[daten.pk]), {"deadline": "15.08.2026"})
    assert resp.status_code == 302
    req = DocumentRequest.objects.get(object=daten)
    assert req.deadline_date == date(2026, 8, 15)
    resp = client.get(reverse("request_detail", args=[req.pk]))
    html = resp.content.decode()
    assert (
        resp.status_code == 200
        and "Version 1" in html
        and "Freigabe erfolgt durch die Geschäftsführung" in html
    )
    resp = client.get(reverse("request_file", args=[req.pk, "pdf"]))
    assert resp.status_code == 200 and resp["Content-Type"].startswith("application/pdf")
    resp = client.post(reverse("request_detail", args=[req.pk]), {"action": "review"})
    req.refresh_from_db()
    assert resp.status_code == 302 and req.status == "reviewed"
    # Freigabe verlangt erneute Anmeldung (Step-up): ohne sie keine Freigabe
    resp = client.post(reverse("request_approve", args=[req.pk]))
    req.refresh_from_db()
    assert resp.status_code == 302 and req.status == "reviewed" and "reauthenticate" in resp["Location"]
