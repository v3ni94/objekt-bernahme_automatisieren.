"""Listen (M11, CR 12a): zwei Laeufe ergeben je Liste und Format genau eine Datei in Drive (zweite Generation
skipped_unchanged, nach Datenaenderung neue Revision derselben ID), Historie mit Alteigentuemer, keine vollstaendige
IBAN in Excel oder PDF (Abbruch mit Fall), Excel-Formatierung, Szenarien geloescht, Papierkorb, verschoben, umbenannt,
Rendering mit langen Werten, Mieterliste, Entprellung."""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import pdfplumber
import pytest
from openpyxl import load_workbook

from apps.audit.models import AuditEvent
from apps.drive.models import DriveNode as DriveNodeRow
from apps.lists import data as data_mod
from apps.lists import services
from apps.lists.models import ListGeneration
from apps.objects.models import ManagedObject, Unit
from apps.parties import services as party_services
from apps.parties.models import Lease, Owner, Tenant, TenantUnitAssignment
from apps.review.models import ReviewCase

pytestmark = pytest.mark.django_db
FULL_IBAN = "DE02120300000000202051"


@pytest.fixture
def daten(welt, fake_oauth, settings, tmp_path):
    settings.OBJEKTAKTE = {**settings.OBJEKTAKTE, "DATA_DIR": tmp_path}
    return welt


def folder(obj, category: str) -> str:
    return DriveNodeRow.objects.get(object=obj, node_kind="main_folder", category_id=category).drive_file_id


def files_in(drive, folder_id: str):
    return [n for n in drive.list_children(folder_id) if not n.is_folder and not n.trashed]


def list_nodes(obj, list_type: str = "owner_list"):
    return DriveNodeRow.objects.filter(
        object=obj, node_kind="list_file", list_type=list_type, status="active"
    ).order_by("list_format")


def own(results):
    return next(r for r in results if r.list_type == "owner_list")


def test_zwei_laeufe_genau_eine_datei_je_format(daten, admin_user):
    obj, drive = daten["objects"]["623"], daten["drive"]
    r1 = services.generate_lists(obj, trigger="manual", user=admin_user, drive=drive)
    assert {r.list_type for r in r1} == set(services.list_types_for(obj))  # 623 ist WEG mit SE-Verwaltung
    assert all(r.status == "done" and r.error is None for r in r1)
    f05 = folder(obj, "05")
    assert sorted(n.name for n in files_in(drive, f05)) == [
        "00_Eigentuemerliste_623.pdf",
        "00_Eigentuemerliste_623.xlsx",
    ]
    nodes = list(list_nodes(obj))
    assert [n.list_format for n in nodes] == ["pdf", "xlsx"] and all(
        n.created_by_app and not n.is_folder for n in nodes
    )
    r2 = services.generate_lists(obj, trigger="run", drive=drive)
    assert own(r2).status == "skipped_unchanged" and len(files_in(drive, f05)) == 2
    assert ListGeneration.objects.filter(object=obj, list_type="owner_list").count() == 4
    assert set(
        ListGeneration.objects.filter(
            object=obj, list_type="owner_list", status="skipped_unchanged"
        ).values_list("drive_node_id", flat=True)
    ) == {n.pk for n in nodes}
    # Datenaenderung: weiterer Eigentuemer -> neue Revision derselben Datei, weiterhin genau eine Datei je Format
    unit = next(u for (num, _label), u in daten["units"].items() if num == "623")
    neu = Owner.objects.create(
        type="natural_person", first_name="Nina", last_name="Neumuster", search_name="NEUMUSTER NINA"
    )
    party_services.create_assignment(
        owner=neu, unit=unit, valid_from=date(2026, 1, 1), valid_to=None, share=Decimal("0.5")
    )
    r3 = services.generate_lists(obj, trigger="masterdata_change", drive=drive)
    assert own(r3).status == "done" and own(r3).content_hash != own(r2).content_hash
    assert len(files_in(drive, f05)) == 2
    for n in nodes:
        assert drive.revisions(n.drive_file_id) == 2 and n.drive_file_id in {
            x.drive_file_id for x in list_nodes(obj)
        }
    assert AuditEvent.objects.filter(action="list.generate", entity_id=obj.pk).count() == 3
    data = data_mod.build_owner_list(obj)
    joint = [r for r in data.current if r["Einheit"] == unit.unit_label]
    assert len(joint) == 2 and all(
        r["Mehrere Eigentümer je Einheit"].startswith("gemeinsam mit") for r in joint
    )


def test_historie_enthaelt_alteigentuemer_und_offene_punkte(daten, admin_user):
    obj = daten["objects"]["623"]
    from apps.requirements.engine import evaluate_object

    evaluate_object(obj, today=date(2026, 7, 1))
    data = data_mod.build_owner_list(obj)
    assert data.history and all(r["Eigentumsende"] is not None for r in data.history)
    ended = {r["Nachname"] for r in data.history}
    assert ended and all(r["Status"] in ("bestätigt", "KI-Vorschlag", "unvollständig") for r in data.current)
    assert data.open_items and set(data.open_columns) == set(data.open_items[0].keys())
    assert all(r["Status"] in ("fehlt", "teilweise") for r in data.open_items)
    services.generate_lists(obj, trigger="manual", user=admin_user, drive=daten["drive"])
    wb = load_workbook(services.local_dir(obj) / "00_Eigentuemerliste_623.xlsx")
    assert wb.sheetnames == ["Aktuell", "Historie", "Offene Punkte"]
    hist = wb["Historie"]
    names = {hist.cell(row=r, column=5).value for r in range(6, hist.max_row + 1)}
    assert ended <= names
    offen = wb["Offene Punkte"]
    assert offen["A5"].value == "Einheit" and offen.max_row >= 6


def test_keine_vollstaendige_iban_in_listen(daten, admin_user):
    obj, drive = daten["objects"]["623"], daten["drive"]
    owner = next(iter(daten["owners"].values()))
    owner.notes = f"Konto {FULL_IBAN}"
    owner.save()
    results = services.generate_lists(obj, trigger="manual", user=admin_user, drive=drive)
    assert own(results).status == "failed" and "Treffer" in own(results).error
    assert ListGeneration.objects.filter(object=obj, list_type="owner_list", status="failed").count() == 2
    assert not list_nodes(obj).exists() and files_in(drive, folder(obj, "05")) == []
    assert not (services.local_dir(obj) / "00_Eigentuemerliste_623.xlsx").exists()
    assert not (services.local_dir(obj) / "00_Eigentuemerliste_623.pdf").exists()
    case = ReviewCase.objects.get(object=obj, case_subtype="list_iban_found")
    assert case.case_type == "move_proposal" and FULL_IBAN not in str(case.context)


def test_excel_formatierung(daten, admin_user):
    obj = daten["objects"]["623"]
    owner = next(iter(daten["owners"].values()))
    owner.correspondence_postal_code, owner.correspondence_city = "01234", "Musterstadt"
    owner.save()
    unit = next(u for (num, _label), u in daten["units"].items() if num == "623")
    unit.house_fee_monthly = Decimal("245.50")
    unit.save()
    services.generate_lists(obj, trigger="manual", user=admin_user, drive=daten["drive"])
    wb = load_workbook(services.local_dir(obj) / "00_Eigentuemerliste_623.xlsx")
    ws = wb["Aktuell"]
    assert (
        ws.freeze_panes == "A6"
        and len(ws.tables) == 1
        and next(iter(ws.tables.values())).ref.startswith("A5:")
    )
    header = [ws.cell(row=5, column=c).value for c in range(1, ws.max_column + 1)]
    assert header[:2] == ["Einheit", "Einheitentyp"] and header[-2:] == ["Status", "Quelle"]
    assert ws["A5"].fill.fgColor.rgb.endswith("E6A83C")
    col = {name: i + 1 for i, name in enumerate(header)}
    plz_cells = [
        ws.cell(row=r, column=col["PLZ"])
        for r in range(6, ws.max_row + 1)
        if ws.cell(row=r, column=col["PLZ"]).value
    ]
    assert plz_cells and all(isinstance(c.value, str) and c.number_format == "@" for c in plz_cells)
    fee = [
        ws.cell(row=r, column=col["Hausgeld monatlich"])
        for r in range(6, ws.max_row + 1)
        if ws.cell(row=r, column=col["Hausgeld monatlich"]).value is not None
    ]
    assert fee and all(c.number_format == "#,##0.00" and isinstance(c.value, float) for c in fee)
    start = [
        ws.cell(row=r, column=col["Eigentumsbeginn"])
        for r in range(6, ws.max_row + 1)
        if ws.cell(row=r, column=col["Eigentumsbeginn"]).value
    ]
    assert start and all(c.number_format == "DD.MM.YYYY" and hasattr(c.value, "year") for c in start)
    assert ws["A1"].value.startswith("Objekt 623") and ws["A3"].value.startswith("Stand: ")
    assert wb.properties.title == "Eigentuemerliste 623"


def test_szenarien_papierkorb_geloescht_verschoben_umbenannt(daten, admin_user):
    obj, drive = daten["objects"]["623"], daten["drive"]
    services.generate_lists(obj, trigger="manual", user=admin_user, drive=drive)
    f05 = folder(obj, "05")
    xlsx, pdf = list_nodes(obj).get(list_format="xlsx"), list_nodes(obj).get(list_format="pdf")
    # Papierkorb: neue Datei, alte Zeile trashed, Papierkorb unveraendert
    drive.trash(xlsx.drive_file_id)
    r = services.generate_lists(obj, trigger="manual", drive=drive)
    assert own(r).status == "done"
    xlsx.refresh_from_db()
    assert (
        xlsx.status == "trashed"
        and list_nodes(obj).get(list_format="xlsx").drive_file_id != xlsx.drive_file_id
    )
    assert len(files_in(drive, f05)) == 2 and drive.get(xlsx.drive_file_id).trashed
    # Endgueltig geloescht (404): neue Datei, alte Zeile missing
    xlsx2 = list_nodes(obj).get(list_format="xlsx")
    drive._entries.pop(xlsx2.drive_file_id)
    services.generate_lists(obj, trigger="manual", drive=drive)
    xlsx2.refresh_from_db()
    assert xlsx2.status == "missing" and len(files_in(drive, f05)) == 2
    # Verschoben und umbenannt: Rueckfuehrung, Inhalt unveraendert (skipped_unchanged)
    root = DriveNodeRow.objects.get(object=obj, node_kind="object_root").drive_file_id
    drive.move(pdf.drive_file_id, f05, root)
    drive.rename(pdf.drive_file_id, "falscher_name.pdf")
    r = services.generate_lists(obj, trigger="manual", drive=drive)
    remote = drive.get(pdf.drive_file_id)
    assert remote.parent_id == f05 and remote.name == "00_Eigentuemerliste_623.pdf"
    assert any("zurück" in m for m in own(r).messages) and drive.revisions(pdf.drive_file_id) == 1
    assert len(files_in(drive, f05)) == 2 and len(files_in(drive, root)) == 0


def test_rendering_lange_werte(daten, admin_user):
    obj = daten["objects"]["623"]
    owner = next(iter(daten["owners"].values()))
    owner.last_name = "Mustermann-Beispielhausen-Altmuster-Neumuster"
    owner.correspondence_street = "Sehr lange Beispielstraße mit vielen Zeichen und Zusätzen"
    owner.correspondence_house_number = "123a"
    owner.notes = "Bemerkung " * 25
    owner.email = "sehr.lange.adresse.mit.vielen.zeichen@beispiel-domain.example"
    owner.save()
    services.generate_lists(obj, trigger="manual", user=admin_user, drive=daten["drive"])
    with pdfplumber.open(str(services.local_dir(obj) / "00_Eigentuemerliste_623.pdf")) as doc:
        text = "\n".join(p.extract_text() or "" for p in doc.pages)
        assert len(doc.pages) >= 1 and doc.pages[0].images
    assert "Eigentümerliste Objekt 623" in text and "Seite 1 von" in text and "HRB 104762" in text
    assert "Aktuell" in text and "Historie" in text and "Offene Punkte" in text
    assert "Neumuster" in text and FULL_IBAN not in text


def test_mieterliste(seeded, admin_user, settings, tmp_path):
    settings.OBJEKTAKTE = {**settings.OBJEKTAKTE, "DATA_DIR": tmp_path}
    obj = ManagedObject.objects.create(
        object_number="700", name="Musterstadt, Beispielweg 7", management_type="rental", is_test=True
    )
    me01 = Unit.objects.create(
        object=obj, unit_label="ME01", unit_label_normalized="ME01", unit_type="apartment"
    )
    Unit.objects.create(
        object=obj,
        unit_label="ME02",
        unit_label_normalized="ME02",
        unit_type="commercial",
        vacancy_confirmed=True,
    )
    tenant = Tenant.objects.create(
        type="natural_person",
        first_name="Anna",
        last_name="Beispiel",
        search_name="BEISPIEL ANNA",
        phone="0000",
        data_status="confirmed",
    )
    lease = Lease.objects.create(
        object=obj,
        start_date=date(2022, 3, 1),
        base_rent=Decimal("650.00"),
        utilities_prepayment=Decimal("120.00"),
        heating_prepayment=Decimal("80.00"),
        deposit_amount=Decimal("1950.00"),
        deposit_type="savings_book",
        rent_adjustment_type="indexed",
        persons_count=2,
        data_status="confirmed",
    )
    TenantUnitAssignment.objects.create(
        tenant=tenant, unit=me01, lease=lease, valid_from=date(2022, 3, 1), data_status="confirmed"
    )
    assert services.list_types_for(obj) == ["tenant_list"]
    data = data_mod.build_tenant_list(obj)
    assert data.columns == data_mod.TENANT_MIN_COLUMNS and len(data.current) == 2
    row = data.current[0]
    assert (
        row["Einheit"] == "ME01"
        and row["Kaltmiete"] == Decimal("650.00")
        and row["Gesamtmiete"] == Decimal("850.00")
    )
    assert (
        row["Kaution Betrag"] == Decimal("1950.00")
        and row["Kaution Anlageform"] == "Sparbuch"
        and row["Staffel oder Index"] == "ja"
    )
    assert row["Anzahl Personen"] == 2 and row["Status"] == "bestätigt" and row["Quelle"] == "manuell"
    leer = data.current[1]
    assert (
        leer["Einheit"] == "ME02" and leer["Bemerkung"] == "Leerstand bestätigt" and leer["Nachname"] is None
    )
    results = services.generate_lists(obj, trigger="manual", user=admin_user, drive=None, publish_files=False)
    assert results[0].status == "done" and (services.local_dir(obj) / "00_Mieterliste_700.xlsx").exists()
    assert ListGeneration.objects.filter(object=obj, drive_node__isnull=True).count() == 2


def test_entprellung_und_job(daten, admin_user, django_capture_on_commit_callbacks, settings):
    settings.OBJEKTAKTE = {**settings.OBJEKTAKTE, "LISTS_AUTO_GENERATE": True}
    obj = daten["objects"]["623"]
    assert services.request_generation(obj.pk, "run") is True  # Automatik eingeschaltet
    services.reset_pending(obj.pk)
    with django_capture_on_commit_callbacks(execute=True):
        assert services.request_generation(obj.pk, "review_confirm") is True
        assert services.request_generation(obj.pk, "review_confirm") is False  # innerhalb der Entprellung
    assert (
        ListGeneration.objects.filter(
            object=obj, list_type="owner_list", trigger_kind="review_confirm"
        ).count()
        == 2
    )
    with django_capture_on_commit_callbacks(execute=True):
        assert services.request_generation(obj.pk, "manual", user_id=admin_user.pk, immediate=True) is True
    assert (
        ListGeneration.objects.filter(
            object=obj, list_type="owner_list", trigger_kind="manual", triggered_by=admin_user
        ).count()
        == 2
    )
    assert list_nodes(obj).count() == 2
