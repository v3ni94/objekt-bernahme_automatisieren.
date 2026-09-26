"""Formatweiche nach Dateiinhalt (26.09.2026): Magic Bytes, Container, E-Mail-Kopfzeilen, HTML. Die Dateien
entstehen im Test; OLE-Container werden nur am Vorspann erkannt, die Stroeme liefert ein Fake (olefile schreibt
keine Container)."""

from __future__ import annotations

import zipfile
from email.message import EmailMessage
from pathlib import Path

import pytest

from apps.pipeline import sniff

OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _datei(tmp_path: Path, name: str, data: bytes) -> Path:
    path = tmp_path / name
    path.write_bytes(data)
    return path


def _eml() -> bytes:
    msg = EmailMessage()
    msg["From"] = "Absender Beispiel <absender@example.test>"
    msg["To"] = "verwaltung@example.test"
    msg["Subject"] = "Testnachricht"
    msg.set_content("Guten Tag")
    return bytes(msg)


def test_pdf_und_bilder(tmp_path):
    assert sniff.detect_kind_by_content(
        _datei(tmp_path, "a", b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    ) == sniff.Sniffed("pdf", ".pdf")
    # Vorspann vor %PDF innerhalb der ersten 1024 Byte ist zulaessig
    assert sniff.detect_kind_by_content(_datei(tmp_path, "b", b"\r\n\r\n%PDF-1.7")).kind == "pdf"
    faelle = {
        "jpg": (b"\xff\xd8\xff\xe0\x00\x10JFIF" + bytes(64), ".jpg"),
        "png": (b"\x89PNG\r\n\x1a\n" + bytes(64), ".png"),
        "tif_le": (b"II*\x00" + bytes(64), ".tif"),
        "tif_be": (b"MM\x00*" + bytes(64), ".tif"),
        "gif": (b"GIF89a" + bytes(64), ".gif"),
        "webp": (b"RIFF\x00\x01\x00\x00WEBPVP8 " + bytes(64), ".webp"),
        "bmp": (b"BM\x36\x04\x00\x00\x00\x00\x00\x00\x36\x00\x00\x00" + bytes(64), ".bmp"),
    }
    for name, (data, suffix) in faelle.items():
        assert sniff.detect_kind_by_content(_datei(tmp_path, name, data)) == sniff.Sniffed("image", suffix), (
            name
        )


def test_zip_container(tmp_path):
    import docx
    import openpyxl

    word = tmp_path / "ohne_endung_docx"
    d = docx.Document()
    d.add_paragraph("Absatz")
    d.save(str(word))
    assert sniff.detect_kind_by_content(word) == sniff.Sniffed("office", ".docx")
    excel = tmp_path / "ohne_endung_xlsx"
    openpyxl.Workbook().save(str(excel))
    assert sniff.detect_kind_by_content(excel) == sniff.Sniffed("office", ".xlsx")
    plain = tmp_path / "archiv"
    with zipfile.ZipFile(plain, "w") as zf:
        zf.writestr("liste.txt", "a;b")
    assert sniff.detect_kind_by_content(plain) == sniff.Sniffed("unsupported", ".zip")
    odt = tmp_path / "text_odf"
    with zipfile.ZipFile(odt, "w") as zf:
        zf.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        zf.writestr("content.xml", "<office:document-content/>")
    assert sniff.detect_kind_by_content(odt) == sniff.Sniffed("office_legacy", ".odt")
    kaputt = _datei(tmp_path, "kaputt", b"PK\x03\x04" + bytes(40))
    assert sniff.detect_kind_by_content(kaputt) == sniff.Sniffed("unsupported", ".zip")


@pytest.mark.parametrize(
    ("streams", "erwartet"),
    [
        (["__properties_version1.0", "__substg1.0_0037001F", "__substg1.0_1000001F"], ("email", ".msg")),
        (["__recip_version1.0_#00000000/__substg1.0_3001001F", "__substg1.0_0037001F"], ("email", ".msg")),
        (["WordDocument", "1Table", "\x05SummaryInformation"], ("office_legacy", ".doc")),
        (["Workbook", "\x05SummaryInformation"], ("office_legacy", ".xls")),
        (["Book"], ("office_legacy", ".xls")),
        (["PowerPoint Document", "Current User"], ("unsupported", ".ppt")),
        ([], ("office_legacy", "")),
    ],
)
def test_ole_container(tmp_path, monkeypatch, streams, erwartet):
    path = _datei(tmp_path, "container", OLE_MAGIC + bytes(1024))
    monkeypatch.setattr(sniff, "_ole_streams", lambda p: streams)
    assert sniff.detect_kind_by_content(path) == sniff.Sniffed(*erwartet)


def test_ole_nicht_lesbar_bleibt_altformat(tmp_path):
    # Vorspann stimmt, Container abgeschnitten: olefile scheitert, Erkennung bleibt beim Altformat ohne Endung
    path = _datei(tmp_path, "abgeschnitten", OLE_MAGIC + bytes(32))
    assert sniff.detect_kind_by_content(path) == sniff.Sniffed("office_legacy", "")


def test_rtf_html_und_email(tmp_path):
    assert sniff.detect_kind_by_content(_datei(tmp_path, "r", b"{\\rtf1\\ansi Text}")) == sniff.Sniffed(
        "office_legacy", ".rtf"
    )
    assert sniff.detect_kind_by_content(
        _datei(tmp_path, "h1", b"<!DOCTYPE html>\n<html><body>x</body></html>")
    ) == sniff.Sniffed("html", ".html")
    assert sniff.detect_kind_by_content(_datei(tmp_path, "h2", b"\xef\xbb\xbf  <HTML><head>")).kind == "html"
    assert (
        sniff.detect_kind_by_content(
            _datei(tmp_path, "h3", b'<?xml version="1.0"?>\n<html xmlns="http://www.w3.org/1999/xhtml">')
        ).kind
        == "html"
    )
    assert sniff.detect_kind_by_content(_datei(tmp_path, "e1", _eml())) == sniff.Sniffed("email", ".eml")
    mbox = b"From absender@example.test Tue Jan 14 09:30:00 2025\n" + _eml()
    assert sniff.detect_kind_by_content(_datei(tmp_path, "e2", mbox)).kind == "email"
    received = (
        b"Received: from mail.example.test\n\tby mx.example.test; Tue, 14 Jan 2025 09:30:00 +0100\n"
        b"Return-Path: <absender@example.test>\nSubject: x\n\nText\n"
    )
    assert sniff.detect_kind_by_content(_datei(tmp_path, "e3", received)).kind == "email"


def test_unbekannt(tmp_path):
    assert sniff.detect_kind_by_content(_datei(tmp_path, "leer", b"")) is None
    assert (
        sniff.detect_kind_by_content(_datei(tmp_path, "text", b"Guten Tag,\ndies ist ein Brief.\n")) is None
    )
    assert sniff.detect_kind_by_content(_datei(tmp_path, "csv", b"Subject;Betrag\nMiete;100\n")) is None
    assert (
        sniff.detect_kind_by_content(_datei(tmp_path, "einzeilig", b"Subject: nur eine Kopfzeile\n")) is None
    )
    assert sniff.detect_kind_by_content(_datei(tmp_path, "binaer", bytes(range(256)))) is None
    assert sniff.detect_kind_by_content(_datei(tmp_path, "xml", b"<?xml version='1.0'?><root/>")) is None
    assert sniff.detect_kind_by_content(tmp_path / "fehlt") is None


def test_temporaere_endungen():
    assert {"", ".tmp", ".herunterladen", ".indir", ".hed", ".part", ".crdownload"} == sniff.TEMP_SUFFIXES


def test_ooxml_hauptteil_entscheidet(tmp_path):
    """Makrofaehige und Vorlagen-Container (.docm, .dotx, .xltx) haben denselben Aufbau wie .docx und .xlsx;
    python-docx und openpyxl lesen sie nicht alle, deshalb entscheidet der Hauptteil-Typ aus [Content_Types].xml."""

    def _paket(name: str, teil: str, ctype: str) -> Path:
        path = tmp_path / name
        with zipfile.ZipFile(path, "w") as zf:
            zf.writestr(
                "[Content_Types].xml",
                '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                '<Default Extension="xml" ContentType="application/xml"/>'
                f'<Override PartName="/{teil}" ContentType="{ctype}"/></Types>',
            )
            zf.writestr(teil, "<x/>")
        return path

    wml = "application/vnd.openxmlformats-officedocument.wordprocessingml"
    sml = "application/vnd.openxmlformats-officedocument.spreadsheetml"
    faelle = {
        "docx": ("word/document.xml", f"{wml}.document.main+xml", ("office", ".docx")),
        "docm": (
            "word/document.xml",
            "application/vnd.ms-word.document.macroEnabled.main+xml",
            ("unsupported", ".docm"),
        ),
        "dotx": ("word/document.xml", f"{wml}.template.main+xml", ("unsupported", ".dotx")),
        "dotm": (
            "word/document.xml",
            "application/vnd.ms-word.template.macroEnabledTemplate.main+xml",
            ("unsupported", ".dotm"),
        ),
        "xlsx": ("xl/workbook.xml", f"{sml}.sheet.main+xml", ("office", ".xlsx")),
        "xlsm": (
            "xl/workbook.xml",
            "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
            ("office", ".xlsm"),
        ),
        "xltx": ("xl/workbook.xml", f"{sml}.template.main+xml", ("unsupported", ".xltx")),
        "pptx": (
            "ppt/presentation.xml",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
            ("unsupported", ".pptx"),
        ),
    }
    for name, (teil, ctype, erwartet) in faelle.items():
        assert sniff.detect_kind_by_content(_paket(name, teil, ctype)) == sniff.Sniffed(*erwartet), name
    # Office-Paket ohne bekannten Hauptteil-Typ: nicht lesbar, kein Endungsvorschlag
    fremd = _paket("fremd", "word/document.xml", "application/x-unbekannt")
    assert sniff.detect_kind_by_content(fremd) == sniff.Sniffed("unsupported", "")


def test_email_mit_abgeschnittener_kopfzeile_am_vorspannende(tmp_path):
    """Exchange-typische ungefaltete Kopfzeilen ueber 8 KB: die Grenze des Vorspanns (HEAD_BYTES) faellt mitten in
    einen Feldnamen; die unvollstaendige Zeile darf nicht als Nicht-Kopfzeile zaehlen."""
    kopf = b"Received: from mail.example.test by mx.example.test\r\nFrom: absender@example.test\r\n"
    kopf += b"X-Microsoft-Antispam-Message-Info: " + b"A" * 7000 + b"\r\n"
    # naechste Zeile so legen, dass Byte 8192 im Feldnamen liegt
    rest = sniff.HEAD_BYTES - len(kopf)
    kopf += b"X-" + b"B" * (rest - 2 + 10) + b": wert\r\n"
    assert len(kopf) > sniff.HEAD_BYTES and kopf[sniff.HEAD_BYTES - 1 : sniff.HEAD_BYTES] == b"B"
    daten = kopf + b"Subject: Testnachricht\r\nTo: verwaltung@example.test\r\n\r\nGuten Tag\r\n"
    assert sniff.detect_kind_by_content(_datei(tmp_path, "lang", daten)) == sniff.Sniffed("email", ".eml")
    # Vorspann endet genau mit einem Zeilenumbruch: nichts wird verworfen
    genau = b"From: a@example.test\r\nTo: b@example.test\r\n"
    genau += b"X-Fuell: " + b"C" * (sniff.HEAD_BYTES - len(genau) - 11) + b"\r\n"
    assert len(genau) == sniff.HEAD_BYTES
    assert sniff.detect_kind_by_content(_datei(tmp_path, "genau", genau + b"\r\nText")).kind == "email"
    # Eine einzige ueberlange Zeile ohne Umbruch bleibt unbekannt
    assert sniff.detect_kind_by_content(_datei(tmp_path, "einzeilig", b"From: " + b"D" * 9000)) is None


def test_name_mit_massgeblicher_endung():
    assert sniff.name_with_suffix("Liste.hed", ".xlsx") == "Liste.xlsx"
    assert sniff.name_with_suffix("Rechnung.tmp", ".pdf") == "Rechnung.pdf"
    assert sniff.name_with_suffix("Rechnung.pdf.tmp", ".pdf") == "Rechnung.pdf"
    assert sniff.name_with_suffix("Rechnung.PDF.crdownload", ".pdf") == "Rechnung.PDF"
    assert sniff.name_with_suffix("Scan", ".pdf") == "Scan.pdf"
    assert sniff.name_with_suffix("Abrechnung.pdf_", ".pdf") == "Abrechnung.pdf"
    assert sniff.name_with_suffix("Bericht.2024", ".pdf") == "Bericht.2024.pdf"
    assert sniff.name_with_suffix("Nachricht.herunterladen", ".msg") == "Nachricht.msg"
    assert sniff.name_with_suffix("Datei.tmp", "") == "Datei.tmp"
    assert (
        sniff.mime_for_suffix(".xlsx") == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert sniff.mime_for_suffix(".pdf") == "application/pdf"
    assert sniff.mime_for_suffix(".eml") == "message/rfc822"
    assert sniff.mime_for_suffix(".unbekannt") == "application/octet-stream"
